"""Excel 导出:各数据集 → xlsx 可读、内容正确(全部离线)。

流水线隔离:segment_stats/周报内部会触 _load_and_process / get_segment_names,
monkeypatch 源模块(函数内 from-import,每次调用重新绑定)。
"""
import os

import openpyxl
import pandas as pd
import pytest

from common.exporter import (
    DATASETS, export_dataset, _funnel_df, _tasks_df,
)
from config.settings import get_settings


@pytest.fixture(autouse=True)
def _no_real_pipeline(monkeypatch):
    """阻止 get_segment_names 触发真实流水线(读 JData + 污染快照)。"""
    import watcher.task_manager as tm_mod
    tm_mod._task_manager = None   # tasks 数据集读单例 → 每测试重置(conftest 已重定向路径)
    seg_df = pd.DataFrame({
        "user_id": [1, 2, 3, 4, 5, 6],
        "recency": [3, 10, 20, 30, 45, 60],
        "frequency": [5.0, 3.0, 2.0, 1.5, 1.0, 1.0],
        "monetary": [800.0, 500.0, 300.0, 200.0, 100.0, 80.0],
        "segment": [2, 2, 1, 1, 0, 0],
        "flow_tag": ["active", "active", "potential", "dormant", "churned", "churned"],
    })
    monkeypatch.setattr(
        "skills.user_segment._load_and_process",
        lambda force_refresh=False: (None, seg_df, None))
    yield
    tm_mod._task_manager = None


def _read_excel(path: str):
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    return wb, ws, rows


class TestExportDatasets:
    def test_tasks_excel_has_header_and_rows(self):
        fn, path = export_dataset("tasks")
        assert fn.endswith(".xlsx") and os.path.exists(path)
        _, ws, rows = _read_excel(path)
        assert [c for c in rows[0]] == ["任务ID", "事件类型", "优先级", "状态", "时间", "摘要"]
        assert ws.max_row >= 1

    def test_funnel_excel_content(self, monkeypatch):
        import pipeline.funnel as pf
        fake = pd.DataFrame({
            "user_id": [1, 1, 2, 2, 2, 3],
            "action_type": [1, 2, 1, 2, 4, 1],
            "date": pd.to_datetime(["2026-08-10 10:00"] * 6),
        })
        monkeypatch.setattr(pf, "load_funnel_actions", lambda force_refresh=False: fake)
        fn, path = export_dataset("funnel", days=7)
        _, ws, rows = _read_excel(path)
        assert ws.max_row - 1 == 3                     # 浏览/加购/下单
        assert rows[0] == ["步骤", "用户数", "转化率", "流失率"]
        assert rows[1][0] == "浏览" and rows[1][1] == 3

    def test_unknown_dataset_raises(self):
        with pytest.raises(ValueError):
            export_dataset("not_a_dataset")

    def test_weekly_report_multi_sheet(self):
        import watcher.weekly_report as wr_mod
        # 幂等防污染:临时目录已由 conftest 重定向;同周文件可能已存在 → force 重新生成
        fn, path = export_dataset("weekly_report")
        wb = openpyxl.load_workbook(path)
        assert len(wb.sheetnames) >= 2                 # 至少"周报"+一节
        assert any("周报" in s for s in wb.sheetnames)

    def test_segment_stats_excel(self):
        fn, path = export_dataset("segment_stats")
        _, ws, rows = _read_excel(path)
        # 合成 seg_df 3 分群 → 3 行数据
        assert ws.max_row - 1 == 3
        header = rows[0]
        assert "segment" in header and "用户数" in header

    def test_segment_trend_empty_snapshots_ok(self):
        # 快照目录为空(conftest 新 tmp)→ 空表不崩、文件可打开
        fn, path = export_dataset("segment_trend")
        _, ws, _ = _read_excel(path)
        assert ws.max_row == 1                         # 仅表头

    def test_segment_growth_insufficient_snapshots(self):
        fn, path = export_dataset("segment_growth")
        _, ws, _ = _read_excel(path)
        assert ws.max_row >= 1                         # 表头存在即可

    def test_segment_users_multi_sheet_per_segment(self):
        fn, path = export_dataset("segment_users")
        wb = openpyxl.load_workbook(path)
        # 合成 seg_df 3 个分群(0/1/2)→ 3 个 sheet
        assert len(wb.sheetnames) == 3
        for sid in (0, 1, 2):
            sheet = next(s for s in wb.sheetnames if s.startswith(f"分群{sid}"))
            ws = wb[sheet]
            assert [c.value for c in ws[1]] == ["用户ID", "近度(天)", "频次", "消费金额", "流转标签"]
            assert ws.max_row - 1 == 2                 # 每分群 2 人
            assert all(isinstance(r[0].value, int) for r in ws.iter_rows(min_row=2))

    def test_segment_users_empty_returns_placeholder(self, monkeypatch):
        import skills.user_segment as us
        monkeypatch.setattr(
            us, "_load_and_process",
            lambda force_refresh=False: (None, pd.DataFrame(columns=["user_id"]), None))
        fn, path = export_dataset("segment_users")
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames == ["数据"]              # 空兜底 sheet


class TestDatasetFrames:
    def test_tasks_df_status_filter(self):
        # _tasks_df 内部用 get_task_manager() 单例 —— autouse fixture 已重置,
        # conftest 已重定向 WATCHER_DB_PATH,直接经单例建任务即可。
        from watcher.task_manager import get_task_manager
        from watcher.events import EventType, Priority, WatcherEvent
        from datetime import datetime, timezone

        tm = get_task_manager()
        for i in range(3):
            tm.create_task(WatcherEvent(
                event_id=f"t{i}", event_type=EventType.SEGMENT_SHIFT,
                priority=Priority.NORMAL, source_snapshot_ts="t", prev_snapshot_ts="",
                details={}, virtual_query=f"事件{i}",
                detected_at=datetime.now(timezone.utc).isoformat(),
            ))
        df = _tasks_df(status="pending")
        assert len(df) == 3
        df_done = _tasks_df(status="completed")
        assert len(df_done) == 0
